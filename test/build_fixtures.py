"""离线生成扩展评测语料；使用 Codex bundled Python 运行，不导入业务应用。"""
from hashlib import sha256
import json
from pathlib import Path
import shutil
from xml.etree import ElementTree as ET
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
DOCS = DATA / "documents"
ASSETS = DATA / "assets"
WORD_NAME = "rag_test_05_星桥实验室操作规程.docx"
PDF_NAME = "rag_test_06_北辰评测方案.pdf"


def snapshot_v1():
    snapshot = DATA / "versions" / "v1"
    if snapshot.exists():
        return
    (snapshot / "documents").mkdir(parents=True)
    shutil.copy2(DATA / "questions.jsonl", snapshot / "questions.jsonl")
    for path in DOCS.glob("*.md"):
        shutil.copy2(path, snapshot / "documents" / path.name)


def extend_md():
    additions = {
        "rag_test_01_青禾课程说明.md": '''
## 学习安排

青禾课程把23天分为三个阶段。第1天至第6天用于环境准备和Python复习，第7天至第16天用于资料处理练习，第17天至第23天用于综合项目和结业复核。每个阶段的学习成果都需要有可复查的文件，不能只提交截图。

环境准备阶段要求学员记录Python版本、包管理器版本和运行命令。环境记录是定位问题的依据，记录时只写版本和目录结构，不写数据库口令或模型密钥。使用虚拟环境时，应明确命令在哪个环境中执行，避免终端与编辑器使用不同解释器。

资料处理阶段包括标题识别、段落组织、代码保护和来源定位。学员需要先阅读原始文档，再观察解析结果。不能因为一份资料能成功上传，就认定所有表格、图片和代码都被完整保留。检查结论应对应到原文件中的具体位置。

综合项目阶段要求使用固定的一组资料演示检索和回答。每次改变切分策略后，要重新生成索引并重新执行问题集。若只修改程序而继续使用旧索引，得到的结果不能说明新策略是否有效。测试时还要记录资料版本，防止不同版本的资料混在一起。

## 学习成果复核

复核首先检查问题是否明确。一个问题如果包含两项独立要求，学员应分别列出需要的证据。例如询问学习周期和作业补交次数时，课程说明可以支持周期，提交规则才能支持补交次数。只返回一份文档并不等于完成全部检索任务。

接着检查事实与对象的对应关系。青禾与白榆是独立课程，两者的周期、最低分数、补交次数和文件格式不能互换。即使回答中包含了所有正确数字，只要把数字归给错误课程，仍然属于错误答案。复核人员需要逐项对照主体和属性。

来源检查要求保留文件名、章节名和原文摘录。原文摘录应足够具体，可以支持相应结论。只引用包含课程名称的介绍段落，不能用来证明该课程的补交次数。引用数量多也不代表证据充分，应看每项关键结论是否都有对应依据。

代码检查要求完整保留命令中的参数。预检查命令与正式执行命令具有不同含义，不能在回答中删除--dry-run等影响行为的参数。代码块中的井号是代码内容，不应当成为新的文档标题。演示时允许把代码作为文本展示，不要求在复核过程中执行代码。

表格检查需要同时识别列标题与行标题。若表格分别列出两门课程的分数，复核人员必须检查数值所在的课程行。不能把同一列的相邻数值合并成一门课程的范围。对于来源不清晰的表格，应返回需要进一步核验，而不是自行补齐关系。

图片检查分为资源保存和图像理解两项。原图可以成功归档，但如果没有执行视觉识别，图中的独有信息未必能够被文本检索找到。复核记录要明确本轮是否启用了视觉理解，不应把未启用视觉能力的结果与启用后的结果直接混为同一组。

多轮对话检查需要使用独立会话。先询问青禾课程，再追问它的结业要求时，系统应当保留青禾这个主体；先询问白榆的补交次数，再问青禾课程呢，系统则应当切换主体。对话历史不能把前一个案例的课程带入下一个独立案例。

无答案检查强调信息边界。资料没有公布收费、教师电话和报名网址时，回答应明确说明缺少资料。即使模型知道类似真实课程的联系方式，也不能把这些常识当作本套虚构资料的事实。复核人员需要确认回答没有生成看似合理的金额、电话或链接。

性能检查应记录首次请求和后续请求的差异。首次请求可能需要加载本地模型，这部分时间应当单独说明。后续请求的中位数与较慢请求的耗时需要同时保留。请求失败时也要记录错误，不能只保留最快的成功样本。

复核最后检查导入是否真正发布。原始文件保存成功与向量写入成功并不是同一个阶段。只有原文、片段、来源和发布凭据彼此一致时，才可以把片段作为知识库证据。遇到资料更新，旧回答所引用的来源需要重新核验。

每次复核应由一人整理问题与证据，再由另一人检查结论。检查人员只评价本轮实际交付的证据，不用后台完整记录中未交付的内容替回答补分。复核表应记录遗漏项、错误归属和无法支持的结论，便于下一轮定位改进效果。

学员提交复核记录后，助教先检查文件是否齐全，再检查案例是否可复现。无法复现的结果应标记为待核验，不能直接宣称通过。复现时保留相同资料、相同问题、相同模型和相同检索参数，只有计划比较的变量可以改变。

跨片段问题要检查头部和尾部信息是否都被保留。一个长小节中的操作要求可能写在前面，而例外处理写在后面。检索时如果只找到了开始部分，应继续查找例外条件。把完整段落切成多个片段后，不能丢弃末尾的有效内容。

版本管理需要区分原始文件内容和系统生成说明。图片摘要属于系统生成内容，不能当作教师原文。学员可以用摘要辅助检索，但关键判断仍要回到图像或原始文本。复核记录要说明证据来自正文、表格还是图片。

复核记录必须填写专用编号，青禾课程的学习成果复核码为QH-R39。该编号仅用于学习成果复核，不是课程编码，也不是学员编号。查找该编号时，应定位到本节末尾，不应从课程编码QH-23A推断或改写。
''',
        "rag_test_02_青禾提交规则.md": '''
## 提交前检查

正式提交前，学员应先解压自己准备的ZIP包，确认README.md位于压缩包根目录，src目录中保留可阅读的源文件。不要把整个虚拟环境或下载缓存打包进入作业，这些文件不会帮助审核人员理解实现过程，反而增加传输与核查成本。

README.md需要说明运行环境、启动方式、输入样本和预期输出。运行环境只记录软件版本，不记录私人账号凭据。启动方式应能在一个新的终端中重现。输入样本尽量采用公开或虚构资料，不能把真实业务中的个人信息放入演示包。

学员应检查中文文件名是否仍然可读。压缩和解压后，目录层级应保持一致，代码中使用的相对路径也应能够找到文件。包含本地图片的文档需要连同图片目录一起打包，不能只上传一份引用了缺失图片的Markdown。

提交前应核对课程编码，青禾作业仍使用QH-23A作为文件名前缀。白榆课程的PDF格式要求不能套用到青禾。若把作业提交到错误课程，系统接受文件并不代表审核人员能在正确的课程列表中找到该记录。

## 反馈与补交

青禾补交规则维持不变：每次收到反馈后的48小时内完成补交，总共最多2次。反馈时间应当来自本次审核记录，不使用截图中没有时区的手工时间。发生时间争议时，保留提交回执和反馈回执，由审核人员核对。

每次补交应解释修改内容，例如修复了路径错误、补齐了来源定位或调整了索引更新流程。补交说明需要对应实际文件变化，不能只写已经解决。审核人员复查时应能够定位到变更的章节或代码位置。

如果反馈指出图片无法显示，先检查原始压缩包是否包含图片，再检查文档中的相对路径。原图缺失与视觉识别失败属于不同问题，前者需要补齐资源，后者需要查看识别配置或重试记录。两者都不能通过编造图片说明解决。

如果反馈指出回答串课程，应使用相同问题重新检索，并检查课程归属和来源文件。仅修改最终回答的措辞不能修复召回范围错误。补交时应提供修改前后的证据列表，让审核人员判断错误是否真正消失。

如果反馈指出切分遗漏代码，需要检查代码围栏的起止位置、完整命令和最后一行。不能通过把代码截图替换正文来绕过完整性检查。补交后的资料仍应同时支持代码阅读和来源核验。

## 回执和归档

青禾作业提交回执的编号前缀为QHACK。收到上传成功提示后还应等待审核任务完成，回执记录文件名、资料版本和提交时间。回执不能替代正式成绩，也不证明内容一定正确。

青禾作业回执保留期为120天。保留期间，学员可以用回执编号查询该次提交的状态。保留期不表示源代码必须在所有设备上保存同样长时间，个人备份由学员自行管理。提交回执中不记录明文密码。

作业更新使用同名替换时，应先保留上一版的评测报告。新版本完成导入后再执行同一组问题，以便比较证据命中和回答完整性。若数据内容也发生变化，报告中要说明这不是单一代码变量的比较。

归档审核需要区分资料包、标准答案和运行结果。资料包是知识库输入，标准答案只供评测使用，运行结果用于解释系统当次行为。把标准答案放进资料包会使测试失去区分检索质量的意义。
''',
        "rag_test_03_白榆课程说明.md": '''
## 零基础学习路径

白榆课程面向零基础学员，因此第一阶段先介绍文件、目录、命令行和简单脚本的含义。学员不需要在报名之前具备Python基础，但学习过程中仍需要完成基础练习。零基础并不表示可以跳过运行环境检查。

课程前5天安排基础工具练习，第6天至第12天安排文档阅读与问题整理，第13天至第17天安排综合演示。白榆的17天学习周期独立于青禾的23天周期，不能把两个时间相加作为任一课程的学习周期。

基础工具练习要求学员记录自己实际输入的命令。遇到错误时，先阅读错误提示，再确认文件是否存在和路径是否正确。未经检查直接重复运行同一命令，通常不能帮助判断问题所在。

文档阅读练习要求学员区分事实、要求与未提供的信息。资料没有写出费用时，应记录为未提供，而不是填写零元。资料没有写出教师电话时，应记录为未知，而不是从其他课程介绍中复制联系方式。

## 成绩构成

白榆课程的结业最低分数仍为75分。白榆平时练习占50%，综合演示占50%。这与青禾平时作业占40%、结业项目占60%的分配不同，比较时必须同时写出课程名称和对应比例。

平时练习检查操作记录是否真实、步骤能否重现、问题描述是否清楚。综合演示检查能否根据资料回答具体问题，以及是否保留必要来源。最终得分由审核记录确定，学员不能通过文件名自行推断成绩。

## PDF作业要求

白榆作业以PDF文件提交，建议使用可搜索的正文文本。图片可以作为说明，但不能把整份作业都转换成低清晰度截图。文档需要有明确标题、步骤说明和结果截图对应的解释。

白榆课程的PDF作业最多12页。页数包括封面和附录，不包括随作业另行保存的原始实验文件。该页数限制只适用于白榆作业，不是PDF解析器的技术页数上限，也不适用于青禾ZIP资料包。

在检查PDF时，应先确认文字可选中和复制，再抽查页码与章节位置。表格需要有完整表头，图片需要清晰且与正文有关。引用的资料来源应保留文件名和章节，不能仅写来自互联网。

白榆作业仍只允许补交1次。补交时应在说明中列出本次修正的内容，并保留初次提交回执。若需要再次调整，应向审核人员申请新的考核安排，不能把青禾课程的2次补交规则套用过来。

## 演示和复核

演示首先说明自己的资料来源，再展示如何从原文找到证据，最后展示回答。若资料不足，应直接说明不足，不应通过展示一段与问题无关的长文来掩盖遗漏。演示记录应能够区分上传成功与回答正确。

复核人员可以选择一条含有两个条件的问题，例如同时询问周期和最低成绩。学员需要确认两个条件都得到回答。若只回答其中一项，记录为部分完成，而不是因为回答语句流畅就认定全部通过。

不同学员使用同一套标准问题时，需要为每个案例建立独立会话。连续追问应在同一会话中执行，但不同案例不能共享未经清理的上下文。这样才能检查系统是否真正理解本题，而不是沿用上一题的主体。
''',
        "rag_test_04_澄石项目手册.md": '''
## 正式导入流程

澄石项目先执行预检查，确认文件格式、资料归属和资源路径，再由操作人员明确执行正式导入。预检查清单包含计划处理的文件和警告，不会凭空生成发布状态。看到预检查通过后，仍需要等待正式导入任务完成。

正式导入首先归档原始文件，再解析正文和图片，之后构建片段、生成向量并写入存储。只有所有必要产物核验成功，才发布该版本的清单。失败任务需要保留可解释的错误信息，不应把半成品当作成功资料。

每批最多64条记录的限制针对写入批次，不是每篇文档只能有64个字或64个段落。较长文档可以生成多个片段，批次应在容量范围内分次提交。命名空间cs_demo_v2用于识别本项目的数据范围，不包含任何凭据信息。

## 配置与清单

澄石项目的正式发布清单文件名为publish_manifest.json。清单必须在向量和来源核验完成后生成，不允许先创建成功清单再补写片段。读取清单失败时，应明确报告依赖不可用，避免把未确认的数据作为证据。

清单记录资料版本、文件指纹和发布的记录编号。相同文件名重传后，旧版本的来源需要重新验证。历史对话可以继续显示先前的文字，但来源下载不能静默转成新版本，从而改变原回答的证据含义。

## 长代码保护案例

下面的脚本是供阅读和切分验证的静态示例。测试工具不运行它，不根据示例创建数据库连接。代码块应整体保留，即使长度超过普通正文片段的目标长度，也不能删除尾部的常量或改变命令参数。

```python
from dataclasses import dataclass
from hashlib import sha256
from pathlib import Path

@dataclass(frozen=True)
class ImportPlan:
    source: Path
    namespace: str = "cs_demo_v2"
    batch_size: int = 64
    dry_run: bool = True

def source_digest(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as stream:
        while True:
            block = stream.read(65536)
            if not block:
                break
            digest.update(block)
    return digest.hexdigest()

def validate_plan(plan: ImportPlan) -> dict:
    if not plan.source.is_file():
        raise ValueError("source file is missing")
    if plan.batch_size < 1 or plan.batch_size > 64:
        raise ValueError("batch size must be between 1 and 64")
    if plan.namespace != "cs_demo_v2":
        raise ValueError("unexpected namespace")
    return {
        "source_file": plan.source.name,
        "source_sha256": source_digest(plan.source),
        "namespace": plan.namespace,
        "batch_size": plan.batch_size,
        "dry_run": plan.dry_run,
    }

def split_batches(records: list, batch_size: int = 64):
    if batch_size < 1 or batch_size > 64:
        raise ValueError("invalid batch size")
    for start in range(0, len(records), batch_size):
        yield records[start:start + batch_size]

def describe_plan(plan: ImportPlan) -> str:
    result = validate_plan(plan)
    fields = [
        f"source={result['source_file']}",
        f"namespace={result['namespace']}",
        f"batch_size={result['batch_size']}",
        f"dry_run={result['dry_run']}",
    ]
    return " | ".join(fields)

def preview_only(plan: ImportPlan) -> dict:
    result = validate_plan(plan)
    result["status"] = "previewed"
    result["published"] = False
    return result

CODE_END_MARKER = "CS-CODE-END-73"
```

代码块最后的校验标记为CS-CODE-END-73。审核时应同时看到函数定义与末尾标记。正文中的解释不能代替检查原始代码是否保留完整，脚本中的井号或字符串也不能被解释为新的Markdown章节。

## 异常与回滚

如果正式导入在发布前失败，先检查当前任务是否已经写出片段，再根据任务边界清理未发布产物。清理只针对本次任务，不能扫描并删除其他课程或其他项目的资料。原始文件可保留用于排查，但不能被标为成功版本。

如果收到向量维度不一致的错误，应检查导入和查询所用的Embedding模型是否一致。改变模型后继续使用旧向量不能得到可靠的比较结果。修复时应明确需要重建哪些资料，不能只修改错误提示文字。

如果返回的来源已经过期，应要求重新检索获取当前版本。系统不应拿同名新文件替换旧回答的来源，也不应在没有确认清单的情况下下载可能已经失效的对象。来源追溯是回答核验的一部分。
''',
    }
    original = DATA / "versions" / "v1" / "documents"
    for name, addition in additions.items():
        text = (original / name).read_text(encoding="utf-8")
        (DOCS / name).write_text(text.rstrip() + "\n\n" + addition.strip() + "\n", encoding="utf-8")


def diagrams():
    from PIL import Image, ImageDraw, ImageFont
    ASSETS.mkdir(exist_ok=True)
    font = ImageFont.truetype("C:/Windows/Fonts/simhei.ttf", 34)
    big = ImageFont.truetype("C:/Windows/Fonts/simhei.ttf", 42)
    image = Image.new("RGB", (1440, 580), "white")
    draw = ImageDraw.Draw(image)
    draw.text((60, 35), "星桥实验室异常处理流程", fill="#142B45", font=big)
    for x, label, time in [(70, "登记", "8分钟"), (535, "复核", "17分钟"), (1000, "归档", "11分钟")]:
        draw.rounded_rectangle((x, 180, x + 370, 430), radius=18, fill="#EAF1F7", outline="#285578", width=3)
        draw.text((x + 120, 220), label, fill="black", font=big)
        draw.text((x + 115, 310), time, fill="#194C6A", font=big)
    for x in [455, 920]:
        draw.line((x, 300, x+60, 300), fill="#285578", width=6)
        draw.polygon([(x+60,300),(x+42,286),(x+42,314)], fill="#285578")
    draw.text((60, 500), "图内时限仅用于异常处理演练", fill="#374151", font=font)
    image.save(ASSETS / "星桥异常处理流程.png")
    image = Image.new("RGB", (1440, 760), "white")
    draw = ImageDraw.Draw(image)
    draw.text((65, 30), "北辰三组吞吐量对比", fill="#142B45", font=big)
    draw.text((65, 95), "单位 每分钟完成请求数", fill="#374151", font=font)
    for x, label, value, color in [(250,"A组",62,"#749BB5"),(650,"B组",86,"#285578"),(1050,"C组",71,"#4F8794")]:
        height=value*5
        draw.rectangle((x,620-height,x+180,620),fill=color)
        draw.text((x+60,620-height-55),str(value),fill="black",font=big)
        draw.text((x+50,650),label,fill="black",font=font)
    draw.line((150,620,1320,620),fill="#52616F",width=3)
    image.save(ASSETS / "北辰吞吐量对比.png")


def create_word():
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.78)
    section.left_margin = section.right_margin = Inches(.83)
    for name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.color.rgb = RGBColor(0,0,0)
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(7)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.22
    doc.add_paragraph("星桥实验室操作规程", "Title")
    doc.add_paragraph("本文说明星桥实验室的资料接收、设备核验、操作记录和异常复核要求。操作人员应先确认样本与记录，再进入演练；来源不明或记录缺失时应暂停处理。")
    doc.add_heading("资料接收与编号", level=1)
    for text in [
        "星桥实验室的接收批次前缀为XQ-LAB，单批最多接收36份样本。接收编号用于追溯资料来源，不是数据库账号，也不能代替样本本身的业务编码。",
        "接收人员逐项检查样本清单、原始文件和提交说明。清单中写出的文件都应能在资料包中找到，名称相近的文件也需要核对内容。接收记录不负责猜测缺失的项目或课程归属。",
        "对于含图片的资料，接收人员先检查图像是否清晰，再检查图像在原文中的位置。图片说明应与相邻正文对应，不得把其他文档的说明复制过来。视觉识别输出需要标明生成来源。",
        "样本登记完成后，复核人员检查编号是否重复。发现重复编号时，先确认是否属于同一资料的新版本，再决定是否替换。不同来源仅仅同名时，需要先更改名称并记录原因，不能直接覆盖。",
        "接收完成不等于实验通过。接收只说明资料满足进入操作环节的条件，后续的环境检查、执行记录和结果核验仍需分别完成。记录中应使用明确的阶段状态，避免把开始处理写成最终成功。",
    ]: doc.add_paragraph(text)
    doc.add_page_break()
    doc.add_heading("设备检查与保留期限", level=1)
    doc.add_paragraph("下表给出星桥实验室三类记录的保留期限。期限从本次操作结束时开始计算，各类记录分别管理。表中的数值不能挪用到青禾作业回执规则。")
    rows = [["记录类型", "保留期限", "用途"], ["设备检查记录", "45天", "确认操作前检查完成"], ["操作日志", "90天", "核对运行过程与异常"], ["归档记录", "180天", "追溯样本与结果版本"]]
    table = doc.add_table(rows=0, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [2.05, 1.25, 3.25]
    for i, values in enumerate(rows):
        cells = table.add_row().cells
        for j, value in enumerate(values):
            cells[j].width = Inches(widths[j])
            cells[j].text = value
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cells[j]._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ["top","left","bottom","right"]:
                item = OxmlElement("w:"+edge)
                for key, val in {"val":"single","sz":"5","color":"D9D9D9"}.items(): item.set(qn("w:"+key),val)
                borders.append(item)
            tcpr.append(borders)
            margins=OxmlElement("w:tcMar")
            for edge in ["top","bottom","left","right"]:
                item=OxmlElement("w:"+edge); item.set(qn("w:w"),"110"); item.set(qn("w:type"),"dxa"); margins.append(item)
            tcpr.append(margins)
            if i==0:
                shade=OxmlElement("w:shd"); shade.set(qn("w:fill"),"E5EDF5"); tcpr.append(shade)
                for run in cells[j].paragraphs[0].runs: run.bold=True
    doc.add_paragraph("检查时应确认设备标识与本次登记一致。检查人员只记录实际看到的状态，不根据上一批次的记录推断当前设备已经准备完毕。异常项需要写明检查位置与现象。")
    doc.add_heading("执行记录", level=1)
    doc.add_paragraph("星桥操作记录必须包含输入版本、执行命令和结果摘要三个字段。只填写运行成功四个字不满足记录要求。对于未完成的步骤，需要说明停止位置，不能补写并未执行的结果。")
    doc.add_paragraph("所有演练命令都使用测试资料。执行记录与标准答案分别保存，标准答案只在评分时使用。若输入资料发生变化，记录新的版本号并重新生成结果，不能沿用旧报告中的成功率。")
    doc.add_page_break()
    doc.add_heading("异常处理流程图", level=1)
    doc.add_paragraph("下图给出异常处理的三个环节与各自时限。时限以图中的文字为准，正文不重复列出。查看流程时应先识别环节，再读取对应时长。")
    doc.add_picture(str(ASSETS / "星桥异常处理流程.png"), width=Inches(6.4))
    doc.add_paragraph("图1 星桥异常处理流程")
    doc.add_heading("复核命令", level=1)
    doc.add_paragraph("星桥实验室使用以下命令校验清单。--check参数表示核验输入清单，不发布演练结果。需要记录完整命令，不能丢弃影响行为的参数。")
    paragraph = doc.add_paragraph()
    run=paragraph.add_run('python verify_lab.py --check --batch XQ-LAB')
    run.font.name="Consolas"; run.font.size=Pt(9)
    # 显式围栏让本地解析器按代码识别，保留完整参数。
    paragraph.text='```bash\npython verify_lab.py --check --batch XQ-LAB\n```'
    doc.add_paragraph("图像时限与文本保留期限属于不同维度。流程图用于指导异常处理节奏，记录表用于说明资料保留时间。回答涉及两者的问题时，应分别引用对应位置，不能把分钟和天互相替换。")
    doc.add_paragraph("星桥实验室的复核签收码为XQ-ACCEPT-58。该签收码只有在清单与操作记录核验完成后才登记，不用于登录任何系统。本文未提供实验室管理员电话、地址或收费信息。")
    # 清除本机 Word 默认模板可能带入的蓝色标题横线。
    for root in [doc.element, doc.styles.element]:
        for border in root.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    doc.save(DOCS / WORD_NAME)


def create_pdf():
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.utils import ImageReader
    font_path="C:/Windows/Fonts/simhei.ttf"
    pdfmetrics.registerFont(TTFont("FixtureCN",font_path))
    width,height=A4
    target=DOCS/PDF_NAME
    pdf=canvas.Canvas(str(target),pagesize=A4)
    pdf.setTitle("北辰评测方案")
    def heading(text,y,size=16):
        pdf.setFont("FixtureCN",size);pdf.setFillColorRGB(0,0,0);pdf.drawString(50,y,text);return y-size-17
    def para(text,y,size=10.5):
        pdf.setFont("FixtureCN",size)
        lines=[];current=""
        for char in text:
            if current and pdfmetrics.stringWidth(current+char,"FixtureCN",size)>width-100:
                lines.append(current);current=""
            current+=char
        if current:lines.append(current)
        for line in lines:
            pdf.drawString(50,y,line);y-=18
        return y-10
    def footer(page):
        pdf.setFont("FixtureCN",9);pdf.setFillColorRGB(.3,.3,.3);pdf.drawRightString(width-50,35,f"北辰评测方案  {page}");pdf.showPage()
    y=heading("北辰评测方案",height-60,23)
    y=para("本文规定北辰评测的样本准备、结果检查和跨文档核验方式。评测人员应先固定资料和问题，再观察检索证据与回答；仅有接口成功状态不能说明答案正确。",y)
    y=heading("样本与分组",y)
    for text in [
        "北辰评测方案的版本编号为BC-EVAL-42，正式问题集包含48个案例。案例分为直接查询、跨文档查询、无答案与多轮追问，分别记录各组结果。",
        "样本准备先确认原始资料可以读取，再核对标准答案与原文摘录。原文、问题和标准答案的职责不同，只有原文进入检索知识库。标准答案只提供给评测人员，不能当作待检索文档上传。",
        "直接查询应要求明确事实，例如版本编号或记录期限。跨文档查询需要标注所有必需证据，不能只标一份最容易找到的资料。无答案问题应检查整个固定语料中都没有可确定的答案。",
        "多轮追问应记录完整上下文和最终问题。评测时让系统真实回答前置问题，再执行追问。不同案例使用独立会话，避免上一题的课程名称影响下一题的检索范围。",
        "北辰评测的复查样本比例为25%。复查人员优先核对跨文档遗漏、主体混淆和来源不支持的结论，再抽查其他案例。复查比例不是合格分数，也不是要求召回的文档比例。",
    ]:y=para(text,y)
    footer(1)
    y=heading("结果判定表",height-60)
    y=para("本页列出北辰评测的三项观察要求。P95耗时使用成功请求统计，同时报告失败数量；首字是第一个实际回答文本片段，进度事件不计作首字。",y)
    rows=[["检查项","本轮观察目标","说明"],["请求成功率","100%","所有计划请求均完成"],["证据全部找齐率","至少90%","适用于有答案案例"],["总耗时P95","不高于15秒","单独记录冷启动"]]
    x=50; col=[125,140,width-100-265];rowheight=42
    for i,row in enumerate(rows):
        offset=x
        for j,txt in enumerate(row):
            pdf.setFillColorRGB(*((.9,.94,.97) if i==0 else (1,1,1)))
            pdf.setStrokeColorRGB(.85,.85,.85);pdf.rect(offset,y-rowheight,col[j],rowheight,fill=1,stroke=1)
            pdf.setFillColorRGB(0,0,0);pdf.setFont("FixtureCN",9.3);pdf.drawString(offset+9,y-25,txt);offset+=col[j]
        y-=rowheight
    y-=28
    for text in [
        "表中的观察目标仅用于这份虚构方案，不是所有RAG系统的行业标准。若改变问题难度、资料规模或硬件环境，应在报告中说明。不同测试集的数值不能脱离条件直接比较。",
        "证据全部找齐率要求每条必需证据都有对应检索片段。一个问题需要两份资料时，只找到其中一份可以算至少命中一条，但不能算证据全部找齐。指标应当同时报告分母和失败案例。",
        "答案判断先检查事实，再检查事实属于哪个对象。把两门课程的成绩要求互换属于错误，即使输出中同时出现两个正确数值。回答增加了资料没有提供的结论，也需要单独记录。",
        "北辰评测报告保留期为210天。报告保存问题版本、原文版本和实际回答，复核表记录评分人员及理由。图像证据应保存关联原图，不能仅保留一段脱离图像的摘要。",
    ]: y=para(text,y)
    footer(2)
    y=heading("吞吐量观察图",height-60)
    y=para("下图比较三个演练组在同一观察窗口内的吞吐量。具体数值仅在图中给出。读取时先核对组名和纵向含义，再回答对应组的数值。",y)
    image=ImageReader(str(ASSETS/"北辰吞吐量对比.png"))
    image_height=(width-100)*760/1440
    pdf.drawImage(image,50,y-image_height,width=width-100,height=image_height)
    y-=image_height+22
    y=para("图1 北辰三组吞吐量对比",y)
    y=heading("图像证据复核",y)
    for text in [
        "图像理解结果需要与原图关联。若本轮未开启视觉理解，可记录图像已归档，但不能宣称已验证图内数值。图像题应单独统计，避免影响纯文本问题的比较。",
        "图片中的组名与正文中的案例类别不同。不能把A组直接解释为直接查询组，也不能把B组解释为跨文档组；图中只给出演练组标签，没有提供这种映射关系。",
        "本文未提供北辰评测服务的访问口令、登录网址或负责人电话。遇到相应问题，应说明资料不足，不能根据版本编号或文件名推断凭据。",
    ]:y=para(text,y)
    footer(3)
    pdf.save()


def write_evidence_sources():
    from pypdf import PdfReader
    sources={}
    for name in [WORD_NAME, PDF_NAME]:
        path=DOCS/name
        if path.suffix==".docx":
            with ZipFile(path) as archive:
                tree=ET.fromstring(archive.read("word/document.xml"))
            ns={"w":"http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            text="\n".join("".join(t.text or "" for t in paragraph.findall(".//w:t",ns)) for paragraph in tree.findall(".//w:p",ns))
        else:
            text="\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
        sources[name]={"sha256":sha256(path.read_bytes()).hexdigest(),"text":text}
    images={
        "星桥异常处理流程.png":"登记8分钟，复核17分钟，归档11分钟。",
        "北辰吞吐量对比.png":"A组62，B组86，C组71，单位为每分钟完成请求数。",
    }
    for name,text in images.items():
        path=ASSETS/name
        sources['assets/'+name]={"sha256":sha256(path.read_bytes()).hexdigest(),"text":text}
    (DATA/"evidence_sources.json").write_text(json.dumps(sources,ensure_ascii=False,indent=2),encoding="utf-8")


def build_questions():
    cases=[json.loads(line) for line in (DATA/"versions/v1/questions.jsonl").read_text(encoding="utf-8").splitlines() if line.strip()]
    def add(number,category,question,answer,evidence,checks,**extra):
        cases.append({"id":f"q{number:02}","category":category,"question":question,"history":[],
                      "answerable":bool(evidence),"expected_answer":answer,"evidence":evidence,
                      "answer_checks":[{"label":label,"any_of":alternatives} for label,alternatives in checks],**extra})
    def e(name,heading,quote,**extra):return {"file":name,"heading":heading,"quote":quote,**extra}
    qh="rag_test_01_青禾课程说明.md";rule="rag_test_02_青禾提交规则.md";by="rag_test_03_白榆课程说明.md";cs="rag_test_04_澄石项目手册.md"
    add(21,"long_section_tail","青禾课程的学习成果复核码是什么？","学习成果复核码是QH-R39。",[e(qh,"学习成果复核","青禾课程的学习成果复核码为QH-R39。")],[("复核码",["QH-R39"])])
    add(22,"long_md","青禾课程三个学习阶段分别安排在哪些天？","第1至6天环境准备和Python复习，第7至16天资料处理，第17至23天综合项目和结业复核。",[e(qh,"学习安排","第1天至第6天用于环境准备和Python复习，第7天至第16天用于资料处理练习，第17天至第23天用于综合项目和结业复核。")],[("阶段一",["第1至6天","1-6天","第1天至第6天"]),("阶段二",["第7至16天","7-16天","第7天至第16天"]),("阶段三",["第17至23天","17-23天","第17天至第23天"])])
    add(23,"multi_evidence","青禾与白榆的平时成绩和最终项目或演示分别占多少？","青禾平时作业40%、结业项目60%；白榆平时练习50%、综合演示50%。",[e(qh,"结业规则","平时作业占 40%，结业项目占 60%。"),e(by,"成绩构成","白榆平时练习占50%，综合演示占50%。")],[("青禾作业",["40%"]),("青禾项目",["60%"]),("白榆练习",["50%"]),("白榆演示",["50%"])])
    add(24,"long_md","青禾补交应在收到反馈后的多少小时内完成？","收到反馈后的48小时内完成。",[e(rule,"补交规则","每次补交必须在收到反馈后的 48 小时内完成。")],[("补交时限",["48小时","四十八小时"])])
    add(25,"multi_evidence","青禾作业回执和北辰评测报告分别保留多久？","青禾作业回执保留120天，北辰评测报告保留210天。",[e(rule,"回执和归档","青禾作业回执保留期为120天。"),e(PDF_NAME,"结果判定表","北辰评测报告保留期为210天。",page=2)],[("青禾回执",["120天"]),("北辰报告",["210天"])])
    add(26,"exact_code_tail","澄石长代码块末尾的CODE_END_MARKER是什么？","CODE_END_MARKER是CS-CODE-END-73。",[e(cs,"长代码保护案例",'CODE_END_MARKER = "CS-CODE-END-73"')],[("代码尾部",["CS-CODE-END-73"])])
    add(27,"word_text","星桥实验室的接收批次前缀与单批样本上限是什么？","批次前缀为XQ-LAB，单批最多36份样本。",[e(WORD_NAME,"资料接收与编号","星桥实验室的接收批次前缀为XQ-LAB，单批最多接收36份样本。")],[("批次前缀",["XQ-LAB"]),("样本上限",["36份","三十六份"])])
    add(28,"word_table","星桥实验室的设备检查记录与归档记录各保留多少天？","设备检查记录保留45天，归档记录保留180天。",[e(WORD_NAME,"设备检查与保留期限","设备检查记录\n45天"),e(WORD_NAME,"设备检查与保留期限","归档记录\n180天")],[("设备记录",["45天"]),("归档记录",["180天"])])
    add(29,"word_table","星桥实验室操作日志的保留期是多少天？","操作日志保留90天。",[e(WORD_NAME,"设备检查与保留期限","操作日志\n90天")],[("操作日志",["90天","九十天"])])
    add(30,"word_code","星桥实验室清单核验使用什么完整命令？","python verify_lab.py --check --batch XQ-LAB",[e(WORD_NAME,"复核命令","python verify_lab.py --check --batch XQ-LAB")],[("完整命令",["python verify_lab.py --check --batch XQ-LAB"])])
    add(31,"word_fields","星桥操作记录必须包含哪三个字段？","输入版本、执行命令和结果摘要。",[e(WORD_NAME,"执行记录","星桥操作记录必须包含输入版本、执行命令和结果摘要三个字段。")],[("输入版本",["输入版本"]),("执行命令",["执行命令"]),("结果摘要",["结果摘要"])])
    add(32,"cross_format","白榆PDF作业最多几页，星桥实验室复核签收码是什么？","白榆PDF作业最多12页，星桥复核签收码为XQ-ACCEPT-58。",[e(by,"PDF作业要求","白榆课程的PDF作业最多12页。"),e(WORD_NAME,"复核命令","星桥实验室的复核签收码为XQ-ACCEPT-58。")],[("白榆页数",["12页"]),("星桥签收码",["XQ-ACCEPT-58"])])
    add(33,"word_image","星桥异常处理流程图中，复核环节的时限是多少分钟？","复核环节的时限为17分钟。",[e(WORD_NAME,"异常处理流程图","复核17分钟",origin="image",asset="assets/星桥异常处理流程.png",match_mode="manual_image")],[("复核时限",["17分钟","十七分钟"])],requires=["vision"])
    add(34,"pdf_text","北辰评测方案的版本编号与正式案例数量是什么？","版本编号BC-EVAL-42，正式问题集48个案例。",[e(PDF_NAME,"样本与分组","北辰评测方案的版本编号为BC-EVAL-42，正式问题集包含48个案例。",page=1)],[("版本编号",["BC-EVAL-42"]),("案例数量",["48个","四十八个"])])
    add(35,"pdf_text","北辰评测的复查样本比例是多少？","复查样本比例为25%。",[e(PDF_NAME,"样本与分组","北辰评测的复查样本比例为25%。",page=1)],[("复查比例",["25%","百分之二十五"])])
    add(36,"pdf_table","北辰方案对证据全部找齐率和总耗时P95的观察目标分别是什么？","证据全部找齐率至少90%，总耗时P95不高于15秒。",[e(PDF_NAME,"结果判定表","证据全部找齐率\n至少90%",page=2),e(PDF_NAME,"结果判定表","总耗时P95\n不高于15秒",page=2)],[("证据目标",["90%","百分之九十"]),("耗时目标",["15秒","十五秒"])])
    add(37,"pdf_false_premise","北辰方案认为HTTP接口成功就说明答案正确，对吗？","不对，仅有接口成功状态不能说明答案正确。",[e(PDF_NAME,"北辰评测方案","仅有接口成功状态不能说明答案正确。",page=1)],[("纠正前提",["不对","不能说明","不代表"])])
    add(38,"cross_format","澄石正式发布清单叫什么，北辰评测报告保留多久？","澄石发布清单为publish_manifest.json，北辰评测报告保留210天。",[e(cs,"配置与清单","澄石项目的正式发布清单文件名为publish_manifest.json。"),e(PDF_NAME,"结果判定表","北辰评测报告保留期为210天。",page=2)],[("清单文件名",["publish_manifest.json"]),("报告期限",["210天"])])
    add(39,"unanswerable","星桥实验室管理员的联系电话是什么？","资料未提供星桥实验室管理员电话，无法确定。",[],[])
    add(40,"pdf_image","北辰吞吐量对比图中，B组每分钟完成多少个请求？","B组每分钟完成86个请求。",[e(PDF_NAME,"吞吐量观察图","B组86",origin="image",asset="assets/北辰吞吐量对比.png",match_mode="manual_image",page=3)],[("B组吞吐",["86个","86次","86","八十六"])],requires=["vision"])
    for case in cases:
        case["source_formats"] = sorted({Path(item["file"]).suffix[1:] for item in case["evidence"]}) or (["docx"] if case["id"] == "q39" else ["md"])
        if case["category"] in {"word_table", "pdf_table"}:
            for item in case["evidence"]:
                item["format"] = "table"
    (DATA/"questions.jsonl").write_text("\n".join(json.dumps(c,ensure_ascii=False) for c in cases)+"\n",encoding="utf-8")


def main():
    snapshot_v1()
    extend_md()
    diagrams()
    create_word()
    create_pdf()
    write_evidence_sources()
    build_questions()
    print(json.dumps({"documents":[{"file":p.name,"bytes":p.stat().st_size} for p in sorted(DOCS.iterdir()) if p.is_file()],"questions":40},ensure_ascii=False,indent=2))


if __name__=="__main__":
    main()
