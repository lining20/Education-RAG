"""DOCX 本地解析：按正文顺序提取标题、段落、列表、表格、代码与图片。"""
from contextlib import contextmanager
from pathlib import Path
from tempfile import TemporaryDirectory
import re
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from common.errors import OperationError
from common.logging.logger import logger
from utils.content_builders.code_blocks import fenced_code
from utils.import_cancellation import check_import_cancelled
from utils.parsers.common import ParseInput, ParseResult, make_asset, make_block, validate_result


def _heading_level(paragraph):
    props = paragraph._p.pPr
    outline = props.find(qn("w:outlineLvl")) if props is not None else None
    if outline is not None:
        value = int(outline.get(qn("w:val")))
        return value + 1 if value < 9 else 0
    style, seen = paragraph.style, set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        match = re.search(r"(?:heading|标题)\s*([1-9])$", style.name or "", re.I)
        if match:
            return int(match[1])
        props = style.element.pPr
        outline = props.find(qn("w:outlineLvl")) if props is not None else None
        if outline is not None:
            value = int(outline.get(qn("w:val")))
            return value + 1 if value < 9 else 0
        style = style.base_style
    return 0


def _code_style(paragraph):
    style, seen = paragraph.style, set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        if any(word in (style.name or "").lower() for word in ("code", "source", "preformatted", "代码")):
            return True
        style = style.base_style
    # 实际课件的代码常使用 Normal + Consolas，而非独立的代码样式。
    monospace = {"consolas", "courier new", "courier", "menlo", "monaco", "cascadia code", "source code pro"}
    runs = [run for run in paragraph.runs if run.text.strip()]
    return bool(runs) and all((run.font.name or (run.style.font.name if run.style else "") or
                              (paragraph.style.font.name if paragraph.style else "") or "").lower() in monospace
                             for run in runs)


@contextmanager
def parse_docx(source: ParseInput):
    """临时图片保持到归档结束，不请求 MinerU，不虚构 Word 页码。"""
    path = source.require_format("docx")
    check_import_cancelled(source.stop_event)
    try:
        document = Document(path)
    except Exception as exc:
        raise OperationError("INVALID_DOCX", "文件无法打开，请检查是否为有效且未加密的 DOCX") from exc
    result, headings, assets = ParseResult(), [], {}
    with TemporaryDirectory(prefix="parse_word_") as directory:
        root = Path(directory)

        def add(kind, text, locator, payload=None):
            result.blocks.append(make_block(source, len(result.blocks), kind, text, headings,
                                            {"format": "docx", **locator}, payload))

        def paragraph_parts(paragraph):
            # 逐个 run 读取，图片之前和之后的文字不会被移动到同一侧。
            parts, text = [], []
            for node in paragraph._p.iter():
                if any(parent.tag == qn("w:txbxContent") for parent in node.iterancestors()):
                    continue
                if node.tag == qn("w:t"):
                    text.append(node.text or "")
                elif node.tag == qn("w:tab"):
                    text.append("\t")
                elif node.tag in {qn("w:br"), qn("w:cr")}:
                    text.append("\n")
                elif node.tag in {qn("a:blip"), "{urn:schemas-microsoft-com:vml}imagedata"}:
                    relation = node.get(qn("r:embed")) or node.get(qn("r:id"))
                    if not relation:
                        raise OperationError("EXTERNAL_WORD_IMAGE", "Word 含外链图片，请先嵌入图片后重新上传")
                    if text:
                        parts.append(("text", "".join(text)))
                        text = []
                    rel = document.part.rels.get(relation)
                    if rel is None or rel.is_external:
                        raise OperationError("INVALID_WORD_IMAGE", "Word 图片关系缺失或指向外部资源")
                    image = rel.target_part
                    identity = str(image.partname)
                    if identity not in assets:
                        target = root / (str(len(assets)) + Path(identity).suffix)
                        target.write_bytes(image.blob)
                        asset = make_asset(source, identity, str(target), identity, image.content_type)
                        assets[identity] = asset
                        result.resources.append(asset)
                    parts.append(("image", assets[identity]))
            if text:
                parts.append(("text", "".join(text)))
            return parts

        pending_code, code_locators, code_images = [], [], []
        fence = None

        def flush_code():
            if pending_code:
                content = "\n".join(pending_code)
                rendered = content if re.match(r"^ {0,3}(`{3,}|~{3,})", content) else fenced_code(content)
                if fence:
                    rendered = content.rstrip("\n") + "\n" + fence
                    result.warnings.append({"code": "WORD_UNCLOSED_CODE_FENCE",
                        "message": "Word 中未闭合的代码围栏按标题、表格或文末边界结束；保留原内容，仅补齐检索文本围栏"})
                if not content.strip():
                    rendered = ""
                add("code" if content.strip() else "image_ref", rendered,
                    {**code_locators[0], "paragraph_end": code_locators[-1]["paragraph_index"]},
                    {"raw_source": rendered, "content": content, "paragraphs": list(code_locators),
                     "images": list(code_images),
                     "resource_ids": list(dict.fromkeys(image["resource_id"] for image in code_images))})
                pending_code.clear()
                code_locators.clear()
                code_images.clear()

        def append_code(text, locator, parts):
            # 图片作为同一代码块的关联资源，记录正文位置，不把说明插进代码。
            offset = sum(len(line) + 1 for line in pending_code)
            for index, (kind, value) in enumerate(parts):
                if kind == "text":
                    offset += len(value)
                else:
                    code_images.append({"resource_id": value.resource_id, **locator,
                                        "inline_index": index, "char_offset": offset,
                                        "char_range_basis": "code_content"})
            pending_code.append(text)
            code_locators.append(locator)

        for ordinal, item in enumerate(document.iter_inner_content()):
            check_import_cancelled(source.stop_event)
            if isinstance(item, Table):
                flush_code()
                fence = None
                rows, images, has_text, code_cells = [], [], False, []
                for row_index, row in enumerate(item.rows):
                    cells, seen = [], set()
                    for col_index, cell in enumerate(row.cells):
                        if cell._tc in seen:
                            continue
                        seen.add(cell._tc)
                        contents, cell_code = [], []
                        for p_index, paragraph in enumerate(cell.paragraphs):
                            parts = paragraph_parts(paragraph)
                            text = "".join(value for kind, value in parts if kind == "text")
                            has_text = has_text or bool(text.strip())
                            if _code_style(paragraph):
                                cell_code.append(text)
                            for kind, value in parts:
                                if kind == "image":
                                    images.append({"resource_id": value.resource_id, "row": row_index, "column": col_index})
                                    contents.append("[图片]")
                                elif not _code_style(paragraph):
                                    contents.append(value)
                        if cell_code:
                            code_cells.append(("\n".join(cell_code), {"table_index": ordinal, "row": row_index,
                                                                   "column": col_index, "paragraph_index": 0}))
                        cells.append("\n".join(contents))
                        if cell.tables:
                            result.warnings.append({"code": "WORD_NESTED_TABLE", "message": "检测到嵌套表格，当前仅提取外层单元格段落，请核对原文件"})
                    rows.append(cells)
                raw = "\n".join("| " + " | ".join(c.replace("|", "\\|").replace("\n", "<br>") for c in row) + " |" for row in rows)
                if has_text or images:
                    add("table" if has_text else "image_ref", raw if has_text else "", {"table_index": ordinal},
                        {"raw_source": raw if has_text else "", "rows": rows, "images": images,
                         "resource_ids": list(dict.fromkeys(i["resource_id"] for i in images))})
                for text, locator in code_cells:
                    rendered = fenced_code(text)
                    add("code", rendered, locator, {"raw_source": rendered, "content": text})
                continue
            locator = {"paragraph_index": ordinal}
            parts = paragraph_parts(item)
            text = "".join(value for kind, value in parts if kind == "text")
            if fence and _heading_level(item):
                flush_code()
                fence = None
            if pending_code and not text.strip() and not parts:
                pending_code.append(text)
                code_locators.append(locator)
                continue
            marker = re.match(r"^ {0,3}(`{3,}|~{3,})(.*)$", text.split("\n", 1)[0])
            if fence or _code_style(item) or marker:
                if not fence and marker:
                    flush_code()
                    fence = marker.group(1)
                elif fence and re.fullmatch(r" {0,3}" + re.escape(fence[0]) + "{" + str(len(fence)) + r",}\s*", text):
                    append_code(text, locator, parts)
                    fence = None
                    flush_code()
                    continue
                append_code(text, locator, parts)
                if fence and any(re.fullmatch(r" {0,3}" + re.escape(fence[0]) + "{" + str(len(fence)) + r",}\s*", line)
                                 for line in text.splitlines()[1:]):
                    fence = None
                    flush_code()
                continue
            flush_code()
            level = _heading_level(item)
            if level and text.strip():
                headings[:] = (headings + [""] * level)[:level - 1] + [text.strip()]
                add("heading", text.strip(), locator, {"level": level})
            for index, (kind, value) in enumerate(parts):
                if kind == "image":
                    add("image_ref", "", {**locator, "inline_index": index}, {"resource_ids": [value.resource_id]})
                elif value.strip() and not level:
                    props = item._p.pPr
                    numbered = (props is not None and props.find(qn("w:numPr")) is not None) or "list" in (item.style.name or "").lower()
                    add("paragraph", value, {**locator, "inline_index": index},
                        {"raw_source": ("- " if numbered else "") + value, "is_list": numbered})
        flush_code()
        if document.element.body.xpath('.//w:txbxContent'):
            result.warnings.append({"code": "WORD_TEXTBOX_UNSUPPORTED", "message": "Word 含文本框，当前未提取其正文；请转 PDF 导入以保留版面内容"})
        if document.element.body.xpath('.//wp:anchor'):
            result.warnings.append({"code": "WORD_FLOATING_IMAGE", "message": "浮动图片按正文锚点关联，位置不代表页面渲染顺序"})
        validate_result(result)
        logger.info("Word 本地解析完成：块数={}，资源数={}", len(result.blocks), len(result.resources))
        yield result
